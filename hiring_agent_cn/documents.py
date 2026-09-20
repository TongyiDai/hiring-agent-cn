"""Document ingestion with visibility checks and evidence provenance."""

from __future__ import annotations

import hashlib
import importlib.util
import os
import tempfile
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pymupdf

from .schemas import DocumentExtraction, EvidenceBlock, SecurityFinding, SecurityReport
from .security import (
    SpanObservation,
    assess_span,
    injection_reasons,
    int_to_rgb,
    intersection_ratio,
    normalize_untrusted_text,
)

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt", ".md", ".png", ".jpg", ".jpeg"}


@dataclass(frozen=True, slots=True)
class DocumentLimits:
    max_file_bytes: int = 25 * 1024 * 1024
    max_pages: int = 20
    max_blocks: int = 5_000
    max_text_characters: int = 250_000
    max_docx_uncompressed_bytes: int = 50 * 1024 * 1024
    max_docx_entries: int = 2_000
    max_image_pixels: int = 40_000_000


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _evidence_id(source_hash: str, page: int | None, index: int, text: str) -> str:
    payload = f"{source_hash}|{page}|{index}|{text}".encode("utf-8", errors="replace")
    return "EV-" + hashlib.sha256(payload).hexdigest()[:12]


def _bbox_tuple(value: object) -> tuple[float, float, float, float] | None:
    try:
        if isinstance(value, pymupdf.Rect):
            return (float(value.x0), float(value.y0), float(value.x1), float(value.y1))
        if isinstance(value, (list, tuple)) and len(value) >= 4:
            return (float(value[0]), float(value[1]), float(value[2]), float(value[3]))
    except (TypeError, ValueError, IndexError):
        return None
    return None


def _trace_text(trace: dict) -> str:
    characters: list[str] = []
    for character in trace.get("chars", []):
        try:
            characters.append(chr(character[0]))
        except (TypeError, ValueError, IndexError):
            continue
    return "".join(characters)


def _page_background_and_coverings(page: pymupdf.Page) -> tuple[list[dict], list[str]]:
    drawings: list[dict] = []
    warnings: list[str] = []
    try:
        drawings = page.get_drawings()
    except Exception as exc:  # pragma: no cover - depends on malformed PDFs
        warnings.append(f"drawing-probe-failed:{type(exc).__name__}")
    return drawings, warnings


def _background_for(
    bbox: tuple[float, float, float, float] | None, seqno: int, drawings: list[dict]
) -> tuple[float, float, float]:
    background = (1.0, 1.0, 1.0)
    latest_seq = -1
    for drawing in drawings:
        fill = drawing.get("fill")
        drawing_bbox = _bbox_tuple(drawing.get("rect"))
        drawing_seq = int(drawing.get("seqno", -1))
        if (
            fill is not None
            and drawing_seq < seqno
            and drawing_seq > latest_seq
            and intersection_ratio(bbox, drawing_bbox) >= 0.8
            and float(drawing.get("fill_opacity") or 1.0) >= 0.8
        ):
            background = (float(fill[0]), float(fill[1]), float(fill[2]))
            latest_seq = drawing_seq
    return background


def _covered_later(
    bbox: tuple[float, float, float, float] | None, seqno: int, page: pymupdf.Page
) -> bool:
    try:
        operations = page.get_bboxlog()
    except Exception:
        return False
    for operation_seq, operation in enumerate(operations):
        if operation_seq <= seqno:
            continue
        kind, operation_bbox = operation[:2]
        if kind not in {"fill-path", "fill-image", "fill-shade"}:
            continue
        if intersection_ratio(bbox, _bbox_tuple(operation_bbox)) >= 0.92:
            return True
    return False


def _outside_page(
    bbox: tuple[float, float, float, float] | None, page_rect: pymupdf.Rect
) -> bool:
    if not bbox:
        return False
    page_bbox = (
        float(page_rect.x0),
        float(page_rect.y0),
        float(page_rect.x1),
        float(page_rect.y1),
    )
    return intersection_ratio(bbox, page_bbox) < 0.5


class PaddleOCRAdapter:
    """Optional OCR adapter. Install the ``ocr`` extra to enable it."""

    def available(self) -> bool:
        return importlib.util.find_spec("paddleocr") is not None

    def extract(
        self, image_path: Path
    ) -> list[tuple[str, float, tuple[float, float, float, float] | None]]:
        if not self.available():
            raise RuntimeError(
                "图片或扫描版简历需要 OCR。请安装 `pip install -e '.[ocr]'`，"
                "或先将文件转换为可复制文字的 PDF/DOCX。"
            )
        from paddleocr import PaddleOCR  # type: ignore

        ocr = PaddleOCR(use_doc_orientation_classify=True, lang="ch")
        result = ocr.predict(str(image_path))
        rows: list[tuple[str, float, tuple[float, float, float, float] | None]] = []
        for page_result in result:
            payload = getattr(page_result, "json", page_result)
            if callable(payload):
                payload = payload()
            payload = payload.get("res", payload) if isinstance(payload, dict) else {}
            texts = payload.get("rec_texts", [])
            scores = payload.get("rec_scores", [])
            boxes = payload.get("rec_boxes", [])
            for text, score, box in zip(texts, scores, boxes, strict=False):
                bbox = _bbox_tuple(box)
                rows.append((str(text), float(score), bbox))
        return rows


class DocumentExtractor:
    def __init__(
        self,
        ocr: PaddleOCRAdapter | None = None,
        limits: DocumentLimits | None = None,
    ):
        self.ocr = ocr or PaddleOCRAdapter()
        self.limits = limits or DocumentLimits()

    def extract(self, path: str | os.PathLike[str]) -> DocumentExtraction:
        source = Path(path).expanduser().resolve()
        if not source.is_file():
            raise FileNotFoundError(f"文件不存在：{source}")
        if source.stat().st_size > self.limits.max_file_bytes:
            raise ValueError(
                f"文件超过 {self.limits.max_file_bytes // (1024 * 1024)}MB 安全上限。"
            )
        suffix = source.suffix.lower()
        if suffix not in SUPPORTED_SUFFIXES:
            raise ValueError(
                f"暂不支持 {suffix or '无扩展名'}，支持：{', '.join(sorted(SUPPORTED_SUFFIXES))}"
            )
        source_hash = file_sha256(source)
        if suffix == ".pdf":
            return self._extract_pdf(source, source_hash)
        if suffix == ".docx":
            return self._extract_docx(source, source_hash)
        if suffix in {".png", ".jpg", ".jpeg"}:
            return self._extract_image(source, source_hash)
        return self._extract_text(source, source_hash)

    def _extract_pdf(self, path: Path, source_hash: str) -> DocumentExtraction:
        blocks: list[EvidenceBlock] = []
        findings: list[SecurityFinding] = []
        warnings: list[str] = []
        probe_warnings: list[str] = []
        with pymupdf.open(path) as document:
            if document.needs_pass:
                raise ValueError("暂不处理加密 PDF，请由候选人提供可审阅版本。")
            if document.page_count > self.limits.max_pages:
                raise ValueError(f"PDF 超过 {self.limits.max_pages} 页安全上限。")
            for page_index, page in enumerate(document):
                page_block_start = len(blocks)
                drawings, drawing_warnings = _page_background_and_coverings(page)
                probe_warnings.extend(
                    f"page-{page_index + 1}:{warning}" for warning in drawing_warnings
                )
                try:
                    traces = page.get_texttrace()
                except Exception as exc:
                    traces = []
                    probe_warnings.append(
                        f"page-{page_index + 1}:texttrace-probe-failed:{type(exc).__name__}"
                    )

                trace_findings: list[
                    tuple[tuple[float, float, float, float] | None, SecurityFinding]
                ] = []
                for trace in traces:
                    text = _trace_text(trace)
                    bbox = _bbox_tuple(trace.get("bbox"))
                    color_value = trace.get("color")
                    color = (
                        (
                            float(color_value[0]),
                            float(color_value[1]),
                            float(color_value[2]),
                        )
                        if isinstance(color_value, (list, tuple)) and len(color_value) >= 3
                        else None
                    )
                    seqno = int(trace.get("seqno", -1))
                    observation = SpanObservation(
                        text=text,
                        page_number=page_index + 1,
                        bbox=bbox,
                        font_size=float(trace.get("size") or 0),
                        color=color,
                        background=_background_for(bbox, seqno, drawings),
                        opacity=float(
                            trace.get("opacity") if trace.get("opacity") is not None else 1.0
                        ),
                        render_mode=int(trace.get("type", 0)),
                        covered_by_later_object=_covered_later(bbox, seqno, page),
                        outside_page=_outside_page(bbox, page.rect),
                    )
                    finding = assess_span(observation)
                    if finding:
                        trace_findings.append((bbox, finding))

                page_dict = page.get_text("dict")
                line_index = 0
                for raw_block in page_dict.get("blocks", []):
                    for line in raw_block.get("lines", []):
                        spans = line.get("spans", [])
                        raw_text = "".join(str(span.get("text") or "") for span in spans)
                        text = normalize_untrusted_text(raw_text)
                        if not text:
                            continue
                        line_bbox = _bbox_tuple(line.get("bbox"))
                        line_findings = [
                            finding
                            for trace_bbox, finding in trace_findings
                            if intersection_ratio(trace_bbox, line_bbox) > 0.15
                            or intersection_ratio(line_bbox, trace_bbox) > 0.15
                        ]
                        direct_observation = SpanObservation(
                            text=raw_text,
                            page_number=page_index + 1,
                            bbox=line_bbox,
                            font_size=min(
                                (float(span.get("size") or 0) for span in spans),
                                default=None,
                            ),
                            color=int_to_rgb(spans[0].get("color")) if spans else None,
                            background=None,
                            opacity=min(
                                (float(span.get("alpha", 255)) / 255 for span in spans),
                                default=None,
                            ),
                            outside_page=_outside_page(line_bbox, page.rect),
                        )
                        direct_finding = assess_span(direct_observation)
                        if direct_finding:
                            line_findings.append(direct_finding)
                        unique = {finding.finding_id: finding for finding in line_findings}
                        if unique:
                            findings.extend(unique.values())
                            continue
                        blocks.append(
                            EvidenceBlock(
                                evidence_id=_evidence_id(
                                    source_hash, page_index + 1, line_index, text
                                ),
                                source_type="pdf",
                                page_number=page_index + 1,
                                bbox=line_bbox,
                                text=text,
                            )
                        )
                        line_index += 1
                        if len(blocks) > self.limits.max_blocks:
                            raise ValueError("文档文本块数量超过安全上限。")
                        if (
                            sum(len(block.text) for block in blocks)
                            > self.limits.max_text_characters
                        ):
                            raise ValueError("文档可见文本长度超过安全上限。")

                if len(blocks) == page_block_start:
                    if self.ocr.available():
                        pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
                        descriptor, raw_path = tempfile.mkstemp(
                            prefix="hiring-agent-cn-page-", suffix=".png"
                        )
                        os.close(descriptor)
                        rendered = Path(raw_path)
                        try:
                            pixmap.save(rendered)
                            for ocr_index, (raw_text, confidence, bbox) in enumerate(
                                self.ocr.extract(rendered)
                            ):
                                text = normalize_untrusted_text(raw_text)
                                if not text:
                                    continue
                                finding = assess_span(
                                    SpanObservation(
                                        text=raw_text, page_number=page_index + 1, bbox=bbox
                                    )
                                )
                                if finding:
                                    findings.append(finding)
                                    continue
                                blocks.append(
                                    EvidenceBlock(
                                        evidence_id=_evidence_id(
                                            source_hash,
                                            page_index + 1,
                                            ocr_index,
                                            text,
                                        ),
                                        source_type="image_ocr",
                                        page_number=page_index + 1,
                                        bbox=bbox,
                                        text=text,
                                        confidence=confidence,
                                        metadata={"origin": "rendered-pdf-page"},
                                    )
                                )
                        finally:
                            rendered.unlink(missing_ok=True)
                    else:
                        warnings.append(
                            f"第 {page_index + 1} 页未提取到可信可见文本；安装 OCR 扩展后可从页面渲染图识别。"
                        )

        if not blocks:
            warnings.append("未取得可用于审阅的文本；扫描版 PDF 请先逐页转换为图片并启用 OCR。")
        return DocumentExtraction(
            source_name=path.name,
            source_sha256=source_hash,
            media_type="application/pdf",
            blocks=blocks,
            safe_text="\n".join(block.text for block in blocks),
            security=SecurityReport(
                findings=list({item.finding_id: item for item in findings}.values()),
                probe_warnings=probe_warnings,
            ),
            warnings=warnings,
        )

    def _extract_docx(self, path: Path, source_hash: str) -> DocumentExtraction:
        from docx import Document

        try:
            with zipfile.ZipFile(path) as archive:
                entries = archive.infolist()
                if len(entries) > self.limits.max_docx_entries:
                    raise ValueError("DOCX 文件条目数量超过安全上限。")
                if (
                    sum(entry.file_size for entry in entries)
                    > self.limits.max_docx_uncompressed_bytes
                ):
                    raise ValueError("DOCX 解压后体积超过安全上限。")
        except zipfile.BadZipFile as exc:
            raise ValueError("DOCX 文件结构无效。") from exc
        document = Document(str(path))
        blocks: list[EvidenceBlock] = []
        findings: list[SecurityFinding] = []
        candidates: list[tuple[str, list[Any]]] = [
            (paragraph.text, paragraph.runs) for paragraph in document.paragraphs
        ]
        for table in document.tables:
            for row in table.rows:
                candidates.append((" | ".join(cell.text for cell in row.cells), []))

        for index, (raw_text, runs) in enumerate(candidates):
            text = normalize_untrusted_text(raw_text)
            if not text:
                continue
            reasons = injection_reasons(text)
            for run in runs:
                if getattr(run.font, "hidden", False):
                    reasons.append("hidden-docx-run")
                if run.font.size and run.font.size.pt < 3:
                    reasons.append("micro-font")
                rgb = getattr(getattr(run.font, "color", None), "rgb", None)
                if rgb and all(
                    int(str(rgb)[offset : offset + 2], 16) > 245 for offset in (0, 2, 4)
                ):
                    reasons.append("low-contrast-text")
            if reasons:
                finding = assess_span(
                    SpanObservation(text=text, font_size=2 if "micro-font" in reasons else 11)
                )
                if not finding:
                    finding = SecurityFinding(
                        finding_id=f"SEC-{hashlib.sha256((str(index) + text).encode()).hexdigest()[:12]}",
                        kind=reasons[0],
                        severity="critical" if injection_reasons(text) else "high",
                        text_sha256=hashlib.sha256(text.encode()).hexdigest(),
                        text_preview=text[:96],
                        reasons=sorted(set(reasons)),
                        action="quarantined",
                    )
                else:
                    finding.reasons = sorted(set(finding.reasons + reasons))
                findings.append(finding)
                continue
            blocks.append(
                EvidenceBlock(
                    evidence_id=_evidence_id(source_hash, None, index, text),
                    source_type="docx",
                    text=text,
                )
            )
            if len(blocks) > self.limits.max_blocks:
                raise ValueError("文档文本块数量超过安全上限。")
        return DocumentExtraction(
            source_name=path.name,
            source_sha256=source_hash,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            blocks=blocks,
            safe_text="\n".join(block.text for block in blocks),
            security=SecurityReport(findings=findings),
        )

    def _extract_text(self, path: Path, source_hash: str) -> DocumentExtraction:
        blocks: list[EvidenceBlock] = []
        findings: list[SecurityFinding] = []
        for index, raw_text in enumerate(path.read_text(encoding="utf-8-sig").splitlines()):
            text = normalize_untrusted_text(raw_text)
            if not text:
                continue
            finding = assess_span(SpanObservation(text=raw_text))
            if finding:
                findings.append(finding)
                continue
            blocks.append(
                EvidenceBlock(
                    evidence_id=_evidence_id(source_hash, None, index, text),
                    source_type="text",
                    text=text,
                )
            )
            if len(blocks) > self.limits.max_blocks:
                raise ValueError("文档文本块数量超过安全上限。")
        return DocumentExtraction(
            source_name=path.name,
            source_sha256=source_hash,
            media_type="text/plain",
            blocks=blocks,
            safe_text="\n".join(block.text for block in blocks),
            security=SecurityReport(findings=findings),
        )

    def _extract_image(self, path: Path, source_hash: str) -> DocumentExtraction:
        blocks: list[EvidenceBlock] = []
        findings: list[SecurityFinding] = []
        try:
            pixmap = pymupdf.Pixmap(str(path))
            if pixmap.width * pixmap.height > self.limits.max_image_pixels:
                raise ValueError("图片像素数量超过安全上限。")
        except RuntimeError as exc:
            raise ValueError("图片文件无法安全读取。") from exc
        for index, (raw_text, confidence, bbox) in enumerate(self.ocr.extract(path)):
            text = normalize_untrusted_text(raw_text)
            if not text:
                continue
            finding = assess_span(SpanObservation(text=raw_text, page_number=1, bbox=bbox))
            if finding:
                findings.append(finding)
                continue
            blocks.append(
                EvidenceBlock(
                    evidence_id=_evidence_id(source_hash, 1, index, text),
                    source_type="image_ocr",
                    page_number=1,
                    bbox=bbox,
                    text=text,
                    confidence=confidence,
                )
            )
        return DocumentExtraction(
            source_name=path.name,
            source_sha256=source_hash,
            media_type=f"image/{path.suffix.lower().lstrip('.')}",
            blocks=blocks,
            safe_text="\n".join(block.text for block in blocks),
            security=SecurityReport(findings=findings),
        )
